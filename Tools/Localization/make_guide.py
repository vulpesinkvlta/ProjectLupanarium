from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
footer = section.footer.paragraphs[0]
footer.text = 'ProjectLupanarium  •  Локализация  •  '
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
footer.style = doc.styles['Normal']; footer.runs[0].font.size = Pt(9)

def inline(p, text):
    for i, part in enumerate(text.split('`')):
        r = p.add_run(part)
        if i % 2:
            r.font.name = 'Consolas'; r.font.size = Pt(9.5)

lines = Path('Docs/Localization/Guide.md').read_text(encoding='utf-8').splitlines()
code = False; table = None
for line in lines:
    if line.startswith('```'):
        code = not code; continue
    if code:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        r = p.add_run(line); r.font.name = 'Consolas'; r.font.size = Pt(9)
        p.paragraph_format.keep_with_next = True
        continue
    if line.startswith('|'):
        if re.match(r'^\|\s*-', line): continue
        cells = [s.strip() for s in line.strip('|').split('|')]
        first = table is None
        if first:
            table = doc.add_table(rows=0, cols=2)
            table.autofit = False
            table.columns[0].width = Inches(2.2); table.columns[1].width = Inches(4.7)
            borders = OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            table._tbl.tblPr.append(borders)
        row = table.add_row()
        for cell, text in zip(row.cells, cells):
            p = cell.paragraphs[0]; p.paragraph_format.space_after=Pt(5); p.paragraph_format.space_before=Pt(5)
            r=p.add_run(text.replace('/', '/\u200b'));r.font.size=Pt(10)
            cell.vertical_alignment=1
            if first:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EBF0');cell._tc.get_or_add_tcPr().append(shade);r.bold=True
        if first:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        continue
    table = None
    if not line: continue
    if line.startswith('# '): doc.add_paragraph(line[2:], 'Title'); continue
    if line.startswith('## '):
        p=doc.add_paragraph(line[3:], 'Heading 1')
        if line[3:] == 'Проверка и сборка': p.paragraph_format.page_break_before=True
        continue
    p=doc.add_paragraph()
    inline(p,line)

doc.core_properties.title='Локализация игры и добавление новых текстов'
doc.core_properties.subject='Русский английский и турецкий языки в ProjectLupanarium'
doc.save('Docs/Localization/LocalizationGuide.docx')
print('Saved Docs/Localization/LocalizationGuide.docx')
